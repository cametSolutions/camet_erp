from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r"C:\Users\T480s\OneDrive\Desktop\camet_erp")
source = root / "output" / "mongodb-api-cost-audit-2026-10-10.md"
out = root / "output" / "mongodb-api-cost-audit-2026-10-10.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def borders(table):
    tbl_pr = table._tbl.tblPr
    node = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9')
        node.append(el)
    tbl_pr.append(node)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7); section.bottom_margin = Inches(.7)
section.left_margin = Inches(.75); section.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9)
for style in ('Heading 1','Heading 2','Heading 3'):
    styles[style].font.name = 'Aptos Display'; styles[style].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('MongoDB API Cost Audit')
r.font.color.rgb = RGBColor(0,0,0)
sub = doc.add_paragraph('Static investigation report for Camet ERP | 10 October 2026')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(10)

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip() or line.startswith('# MongoDB'):
        i += 1; continue
    if line.startswith('## '):
        doc.add_heading(line[3:], level=1); i += 1; continue
    if line.startswith('### '):
        doc.add_heading(line[4:], level=2); i += 1; continue
    if line.startswith('|'):
        rows=[]
        while i < len(lines) and lines[i].startswith('|'):
            if not set(lines[i].replace('|','').strip()) <= set('-: '): rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i += 1
        if rows:
            table=doc.add_table(rows=1, cols=len(rows[0])); table.style='Table Grid'; borders(table)
            for c, text in zip(table.rows[0].cells, rows[0]):
                c.text=text; shade(c,'1F4E78')
                for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(8)
            for ridx,row in enumerate(rows[1:]):
                cells=table.add_row().cells
                for c,text in zip(cells,row):
                    c.text=text
                    if ridx % 2 == 1: shade(c,'F2F6FA')
                    for p in c.paragraphs:
                        for run in p.runs: run.font.size=Pt(7.5)
            doc.add_paragraph()
        continue
    if line.startswith('- '):
        doc.add_paragraph(line[2:], style='List Bullet'); i += 1; continue
    if line[:2].isdigit() and '. ' in line[:5]:
        doc.add_paragraph(line.split('. ',1)[1], style='List Number'); i += 1; continue
    p=doc.add_paragraph()
    p.add_run(line.replace('`',''))
    i += 1

doc.save(out)
print(out)
